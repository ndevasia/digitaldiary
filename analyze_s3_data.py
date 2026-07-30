#!/usr/bin/env python3
"""
Analyze downloaded S3 Digital Diary data:
- Count media files (screenshots, videos, audio)
- Calculate total duration of media
- Parse logs for feature usage
- Generate timeline visualizations
"""

import json
import subprocess
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from typing import Dict, List, Tuple
import statistics

# Dependencies: pip install matplotlib pandas openpyxl
try:
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    import pandas as pd
    HAS_PLOTTING = True
except ImportError:
    HAS_PLOTTING = False
    print("⚠️  Optional: Install matplotlib and pandas for visualizations: pip install matplotlib pandas openpyxl")

DATA_DIR = Path("s3_data")


class MediaAnalyzer:
    def __init__(self, data_dir: Path = DATA_DIR):
        self.data_dir = data_dir
        self.participants = [d.name for d in data_dir.iterdir() if d.is_dir()]
        self.results = {}

    def get_media_counts(self) -> Dict[str, Dict[str, int]]:
        """Count screenshots, videos, and audio files per participant."""
        counts = {}
        
        for participant in self.participants:
            user_dir = self.data_dir / participant
            media_dir = user_dir / "media"
            
            counts[participant] = {
                "screenshots": 0,
                "screen_recordings": 0,
                "audio_recordings": 0,
                "total": 0
            }
            
            if not media_dir.exists():
                continue
            
            # Count screenshots
            screenshot_dir = media_dir / "screenshots"
            if screenshot_dir.exists():
                counts[participant]["screenshots"] = len(list(screenshot_dir.glob("*")))
            
            # Count screen recordings
            recording_dir = media_dir / "screen_recordings"
            if recording_dir.exists():
                counts[participant]["screen_recordings"] = len(list(recording_dir.glob("*")))
            
            # Count audio recordings
            audio_dir = media_dir / "audio_recordings"
            if audio_dir.exists():
                counts[participant]["audio_recordings"] = len(list(audio_dir.glob("*")))
            
            counts[participant]["total"] = sum([
                counts[participant]["screenshots"],
                counts[participant]["screen_recordings"],
                counts[participant]["audio_recordings"]
            ])
        
        return counts

    def get_media_duration(self) -> Dict[str, Dict[str, float]]:
        """Calculate total duration (in minutes) of video/audio files per participant."""
        durations = {}
        
        for participant in self.participants:
            user_dir = self.data_dir / participant
            media_dir = user_dir / "media"
            
            durations[participant] = {
                "screen_recordings_minutes": 0.0,
                "audio_recordings_minutes": 0.0,
                "total_minutes": 0.0
            }
            
            if not media_dir.exists():
                continue
            
            # Get recording durations
            recording_dir = media_dir / "screen_recordings"
            if recording_dir.exists():
                total = self._get_directory_duration(recording_dir)
                durations[participant]["screen_recordings_minutes"] = total
            
            # Get audio durations
            audio_dir = media_dir / "audio_recordings"
            if audio_dir.exists():
                total = self._get_directory_duration(audio_dir)
                durations[participant]["audio_recordings_minutes"] = total
            
            durations[participant]["total_minutes"] = (
                durations[participant]["screen_recordings_minutes"] +
                durations[participant]["audio_recordings_minutes"]
            )
        
        return durations

    def _get_directory_duration(self, directory: Path) -> float:
        """Get total duration of all media files in a directory (in minutes)."""
        total_seconds = 0.0
        
        for file_path in directory.glob("*"):
            if not file_path.is_file():
                continue
            
            try:
                duration = self._get_file_duration(file_path)
                if duration:
                    total_seconds += duration
            except Exception as e:
                print(f"  ⚠️  Could not get duration for {file_path.name}: {e}")
        
        return total_seconds / 60.0 if total_seconds > 0 else 0.0

    def _get_file_duration(self, file_path: Path) -> float:
        """Get duration of a media file in seconds using ffprobe."""
        try:
            result = subprocess.run(
                [
                    "ffprobe",
                    "-v", "error",
                    "-show_entries", "format=duration",
                    "-of", "default=noprint_wrappers=1:nokey=1:noprint_wrappers=1",
                    str(file_path)
                ],
                capture_output=True,
                text=True,
                timeout=5
            )
            
            if result.stdout.strip():
                return float(result.stdout.strip())
        except (subprocess.TimeoutExpired, FileNotFoundError):
            pass
        
        return None

    def parse_logs(self) -> Dict[str, Dict[str, int]]:
        """Parse logs to count feature usage per participant."""
        feature_counts = defaultdict(lambda: defaultdict(int))
        
        for participant in self.participants:
            user_dir = self.data_dir / participant
            logs_dir = user_dir / "logs"
            
            if not logs_dir.exists():
                continue
            
            for log_file in logs_dir.glob("*.json"):
                try:
                    with open(log_file, 'r') as f:
                        logs = json.load(f)
                        if isinstance(logs, list):
                            for log in logs:
                                if "feature" in log:
                                    feature = log["feature"]
                                    feature_counts[participant][feature] += 1
                except json.JSONDecodeError:
                    print(f"  ⚠️  Could not parse {log_file.name}")
        
        return dict(feature_counts)

    def parse_page_sources(self) -> Dict[str, Dict[str, int]]:
        """Parse logs to count page visits per participant."""
        page_counts = defaultdict(lambda: defaultdict(int))
        
        for participant in self.participants:
            user_dir = self.data_dir / participant
            logs_dir = user_dir / "logs"
            
            if not logs_dir.exists():
                continue
            
            for log_file in logs_dir.glob("*.json"):
                try:
                    with open(log_file, 'r') as f:
                        logs = json.load(f)
                        if isinstance(logs, list):
                            for log in logs:
                                if "page_source" in log:
                                    page = log["page_source"]
                                    page_counts[participant][page] += 1
                except json.JSONDecodeError:
                    pass
        
        return dict(page_counts)

    def get_timeline_data(self) -> Dict[str, List[Tuple[datetime, str]]]:
        """Extract timestamps and features from logs for timeline visualization."""
        timeline = defaultdict(list)
        
        for participant in self.participants:
            user_dir = self.data_dir / participant
            logs_dir = user_dir / "logs"
            
            if not logs_dir.exists():
                continue
            
            for log_file in logs_dir.glob("*.json"):
                try:
                    with open(log_file, 'r') as f:
                        logs = json.load(f)
                        if isinstance(logs, list):
                            for log in logs:
                                if "timestamp" in log and "feature" in log:
                                    try:
                                        ts = datetime.fromisoformat(log["timestamp"])
                                        timeline[participant].append((ts, log["feature"]))
                                    except ValueError:
                                        pass
                except json.JSONDecodeError:
                    pass
        
        return dict(timeline)

    def print_summary(self):
        """Print analysis summary."""
        print("\n" + "="*80)
        print("DIGITAL DIARY DATA ANALYSIS SUMMARY")
        print("="*80)
        
        # Media counts
        print("\n📊 MEDIA COUNTS PER PARTICIPANT")
        print("-" * 80)
        
        media_counts = self.get_media_counts()
        total_screenshots = 0
        total_videos = 0
        total_audio = 0
        
        for participant in sorted(self.participants):
            counts = media_counts[participant]
            total_screenshots += counts["screenshots"]
            total_videos += counts["screen_recordings"]
            total_audio += counts["audio_recordings"]
            
            print(f"\n{participant:15} | Screenshots: {counts['screenshots']:3d} | "
                  f"Videos: {counts['screen_recordings']:3d} | Audio: {counts['audio_recordings']:3d} | "
                  f"Total: {counts['total']:3d}")
        
        print(f"\n{'TOTAL':15} | Screenshots: {total_screenshots:3d} | "
              f"Videos: {total_videos:3d} | Audio: {total_audio:3d}")
        
        # Media duration
        print("\n\n⏱️  MEDIA DURATION PER PARTICIPANT (minutes)")
        print("-" * 80)
        
        print("Note: Calculating duration... (requires ffprobe)")
        durations = self.get_media_duration()
        
        total_video_mins = 0
        total_audio_mins = 0
        
        for participant in sorted(self.participants):
            dur = durations[participant]
            total_video_mins += dur["screen_recordings_minutes"]
            total_audio_mins += dur["audio_recordings_minutes"]
            
            if dur["total_minutes"] > 0:
                print(f"\n{participant:15} | Videos: {dur['screen_recordings_minutes']:8.1f} min | "
                      f"Audio: {dur['audio_recordings_minutes']:8.1f} min | "
                      f"Total: {dur['total_minutes']:8.1f} min")
        
        if total_video_mins > 0 or total_audio_mins > 0:
            print(f"\n{'TOTAL':15} | Videos: {total_video_mins:8.1f} min | "
                  f"Audio: {total_audio_mins:8.1f} min | "
                  f"Total: {total_video_mins + total_audio_mins:8.1f} min")
        
        # Feature usage
        print("\n\n🎯 FEATURE USAGE PER PARTICIPANT")
        print("-" * 80)
        
        feature_counts = self.parse_logs()
        
        for participant in sorted(self.participants):
            features = feature_counts.get(participant, {})
            if not features:
                print(f"\n{participant:15} | No log data")
                continue
            
            print(f"\n{participant:15} |")
            
            # Sort by frequency
            sorted_features = sorted(features.items(), key=lambda x: x[1], reverse=True)
            
            for feature, count in sorted_features[:5]:  # Top 5
                print(f"  {feature:35s} : {count:4d} calls")
            
            if len(sorted_features) > 5:
                print(f"  ... and {len(sorted_features) - 5} more features")

    def print_page_visits(self):
        """Print page visit summary."""
        print("\n\n📄 PAGE VISITS PER PARTICIPANT")
        print("-" * 80)
        
        page_counts = self.parse_page_sources()
        
        for participant in sorted(self.participants):
            pages = page_counts.get(participant, {})
            if not pages:
                print(f"\n{participant:15} | No log data")
                continue
            
            print(f"\n{participant:15} |")
            
            # Sort by frequency
            sorted_pages = sorted(pages.items(), key=lambda x: x[1], reverse=True)
            
            for page, count in sorted_pages:
                print(f"  {page:35s} : {count:4d} visits")

    def create_visualizations(self):
        """Create matplotlib visualizations."""
        if not HAS_PLOTTING:
            print("⚠️  Skipping visualizations. Install matplotlib: pip install matplotlib pandas")
            return
        
        print("\n\n📈 CREATING VISUALIZATIONS...")
        print("-" * 80)
        
        # Get data
        media_counts = self.get_media_counts()
        feature_counts = self.parse_logs()
        timeline_data = self.get_timeline_data()
        
        # 1. Media counts by type
        fig, axes = plt.subplots(2, 2, figsize=(16, 12))
        fig.suptitle("Digital Diary Analysis", fontsize=16, fontweight='bold')
        
        # Subplot 1: Media counts stacked bar
        participants = sorted(self.participants)
        screenshots = [media_counts[p]["screenshots"] for p in participants]
        videos = [media_counts[p]["screen_recordings"] for p in participants]
        audio = [media_counts[p]["audio_recordings"] for p in participants]
        
        x = range(len(participants))
        ax = axes[0, 0]
        ax.bar(x, screenshots, label='Screenshots', color='#3498db')
        ax.bar(x, videos, bottom=screenshots, label='Videos', color='#2ecc71')
        ax.bar(x, audio, bottom=[s+v for s,v in zip(screenshots, videos)], label='Audio', color='#e74c3c')
        ax.set_ylabel('Count')
        ax.set_title('Media Files per Participant')
        ax.set_xticks(x)
        ax.set_xticklabels(participants, rotation=45, ha='right')
        ax.legend()
        ax.grid(axis='y', alpha=0.3)
        
        # Subplot 2: Total media count
        totals = [media_counts[p]["total"] for p in participants]
        ax = axes[0, 1]
        ax.barh(participants, totals, color='#9b59b6')
        ax.set_xlabel('Total Media Count')
        ax.set_title('Total Media Files per Participant')
        ax.grid(axis='x', alpha=0.3)
        
        # Subplot 3: Top features
        ax = axes[1, 0]
        all_features = defaultdict(int)
        for participant, features in feature_counts.items():
            for feature, count in features.items():
                all_features[feature] += count
        
        top_features = sorted(all_features.items(), key=lambda x: x[1], reverse=True)[:10]
        if top_features:
            features_names = [f[0][:30] for f in top_features]  # Truncate names
            features_counts = [f[1] for f in top_features]
            ax.barh(features_names, features_counts, color='#f39c12')
            ax.set_xlabel('Usage Count')
            ax.set_title('Top 10 Most Used Features')
            ax.grid(axis='x', alpha=0.3)
        
        # Subplot 4: Participants by feature usage intensity
        total_calls = {p: sum(features.values()) for p, features in feature_counts.items()}
        ax = axes[1, 1]
        sorted_by_calls = sorted(total_calls.items(), key=lambda x: x[1], reverse=True)
        if sorted_by_calls:
            p_names = [p[0] for p in sorted_by_calls]
            p_calls = [p[1] for p in sorted_by_calls]
            ax.bar(range(len(p_names)), p_calls, color='#1abc9c')
            ax.set_ylabel('Total API Calls')
            ax.set_title('Feature Usage Intensity by Participant')
            ax.set_xticks(range(len(p_names)))
            ax.set_xticklabels(p_names, rotation=45, ha='right')
            ax.grid(axis='y', alpha=0.3)
        
        plt.tight_layout()
        output_file = Path("analysis_overview.png")
        plt.savefig(output_file, dpi=150, bbox_inches='tight')
        print(f"✅ Saved: {output_file}")
        
        # 2. Timeline visualization
        if timeline_data and any(timeline_data.values()):
            fig, ax = plt.subplots(figsize=(16, 8))
            
            for participant, events in timeline_data.items():
                if not events:
                    continue
                
                events.sort(key=lambda x: x[0])
                timestamps = [e[0] for e in events]
                
                # Count events per day
                daily_counts = defaultdict(int)
                for ts in timestamps:
                    day = ts.date()
                    daily_counts[day] += 1
                
                days = sorted(daily_counts.keys())
                counts = [daily_counts[d] for d in days]
                
                ax.plot(days, counts, marker='o', label=participant, alpha=0.7)
            
            ax.set_xlabel('Date')
            ax.set_ylabel('Daily API Calls')
            ax.set_title('Feature Usage Timeline')
            ax.legend(bbox_to_anchor=(1.05, 1), loc='upper left')
            ax.grid(True, alpha=0.3)
            ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y-%m-%d'))
            plt.setp(ax.xaxis.get_majorticklabels(), rotation=45)
            
            plt.tight_layout()
            output_file = Path("analysis_timeline.png")
            plt.savefig(output_file, dpi=150, bbox_inches='tight')
            print(f"✅ Saved: {output_file}")
        
        plt.close('all')

    def export_to_excel(self):
        """Export analysis to Excel spreadsheet."""
        if not HAS_PLOTTING:
            print("⚠️  Skipping Excel export. Install openpyxl: pip install openpyxl")
            return
        
        print("\n📋 CREATING EXCEL REPORT...")
        print("-" * 80)
        
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            print("⚠️  Install openpyxl: pip install openpyxl")
            return
        
        # Get data
        media_counts = self.get_media_counts()
        durations = self.get_media_duration()
        feature_counts = self.parse_logs()
        
        # Create workbook
        wb = openpyxl.Workbook()
        
        # Sheet 1: Media Counts
        ws = wb.active
        ws.title = "Media Counts"
        
        ws['A1'] = "Participant"
        ws['B1'] = "Screenshots"
        ws['C1'] = "Screen Recordings"
        ws['D1'] = "Audio Recordings"
        ws['E1'] = "Total"
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            ws[f'{col}1'].font = Font(bold=True)
            ws[f'{col}1'].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            ws[f'{col}1'].font = Font(bold=True, color="FFFFFF")
        
        row = 2
        for participant in sorted(self.participants):
            counts = media_counts[participant]
            ws[f'A{row}'] = participant
            ws[f'B{row}'] = counts["screenshots"]
            ws[f'C{row}'] = counts["screen_recordings"]
            ws[f'D{row}'] = counts["audio_recordings"]
            ws[f'E{row}'] = counts["total"]
            row += 1
        
        # Sheet 2: Duration
        ws = wb.create_sheet("Duration (Minutes)")
        
        ws['A1'] = "Participant"
        ws['B1'] = "Screen Recording (min)"
        ws['C1'] = "Audio Recording (min)"
        ws['D1'] = "Total (min)"
        
        for col in ['A', 'B', 'C', 'D']:
            ws[f'{col}1'].font = Font(bold=True, color="FFFFFF")
            ws[f'{col}1'].fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        
        row = 2
        for participant in sorted(self.participants):
            dur = durations[participant]
            ws[f'A{row}'] = participant
            ws[f'B{row}'] = f"{dur['screen_recordings_minutes']:.1f}"
            ws[f'C{row}'] = f"{dur['audio_recordings_minutes']:.1f}"
            ws[f'D{row}'] = f"{dur['total_minutes']:.1f}"
            row += 1
        
        # Sheet 3: Feature Usage
        ws = wb.create_sheet("Feature Usage")
        
        ws['A1'] = "Participant"
        ws['B1'] = "Feature"
        ws['C1'] = "Count"
        
        for col in ['A', 'B', 'C']:
            ws[f'{col}1'].font = Font(bold=True, color="FFFFFF")
            ws[f'{col}1'].fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        
        row = 2
        for participant in sorted(self.participants):
            features = feature_counts.get(participant, {})
            sorted_features = sorted(features.items(), key=lambda x: x[1], reverse=True)[:5]
            
            first_row_for_participant = True
            for feature, count in sorted_features:
                ws[f'A{row}'] = participant if first_row_for_participant else ""
                ws[f'B{row}'] = feature
                ws[f'C{row}'] = count
                first_row_for_participant = False
                row += 1
        
        # Sheet 4: Page Visits
        ws = wb.create_sheet("Page Visits")
        
        page_counts = self.parse_page_sources()
        
        ws['A1'] = "Participant"
        ws['B1'] = "Page"
        ws['C1'] = "Visits"
        
        for col in ['A', 'B', 'C']:
            ws[f'{col}1'].font = Font(bold=True, color="FFFFFF")
            ws[f'{col}1'].fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        
        row = 2
        for participant in sorted(self.participants):
            pages = page_counts.get(participant, {})
            sorted_pages = sorted(pages.items(), key=lambda x: x[1], reverse=True)[:5]
            
            first_row_for_participant = True
            for page, count in sorted_pages:
                ws[f'A{row}'] = participant if first_row_for_participant else ""
                ws[f'B{row}'] = page
                ws[f'C{row}'] = count
                first_row_for_participant = False
                row += 1
        
        # Auto-adjust column widths
        for ws in wb.sheetnames:
            for column in wb[ws].columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                wb[ws].column_dimensions[column_letter].width = adjusted_width
        
        output_file = Path("analysis_report.xlsx")
        wb.save(output_file)
        print(f"✅ Saved: {output_file}")


def main():
    print("🔍 Analyzing Digital Diary S3 Data...")
    
    if not DATA_DIR.exists():
        print(f"❌ Data directory not found: {DATA_DIR}")
        print("Run 'python download_s3_data.py' first to download the data.")
        return
    
    analyzer = MediaAnalyzer()
    
    if not analyzer.participants:
        print(f"❌ No participant data found in {DATA_DIR}")
        return
    
    print(f"📁 Found {len(analyzer.participants)} participants")
    
    analyzer.print_summary()
    analyzer.print_page_visits()
    analyzer.create_visualizations()
    analyzer.export_to_excel()
    
    print("\n" + "="*80)
    print("✅ Analysis complete!")
    print("="*80)


if __name__ == "__main__":
    main()
